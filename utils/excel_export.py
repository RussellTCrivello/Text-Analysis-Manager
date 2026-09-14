"""
Excel Export Utility
Professional export with centralized header settings
Enhanced with full Arabic RTL support
"""
from typing import List, Dict, Optional
from datetime import datetime
from translations.translations import TranslationManager
from utils.logger import get_logger

# Import print settings
try:
    from utils.print_utils import PrintSettings, get_print_settings
except ImportError:
    PrintSettings = None
    get_print_settings = None

logger = get_logger(__name__)


def get_arabic_font_name():
    """Get the appropriate Arabic font name"""
    import os
    fonts_dir = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'fonts')
    noto_path = os.path.join(fonts_dir, 'NotoSansArabic-Regular.ttf')
    
    if os.path.exists(noto_path):
        return 'Noto Sans Arabic'
    return 'Traditional Arabic'


def export_to_excel(
    data: List[Dict],
    columns: List[tuple],
    filename: str,
    translator: TranslationManager,
    table_name: str = "Data",
    include_all_fields: bool = True,
    selected_columns: Optional[List[tuple]] = None,
    column_widths: Optional[Dict[str, int]] = None
) -> bool:
    """
    Export data to Excel file with professional header and formatting
    
    Args:
        data: List of dictionaries containing data to export
        columns: List of tuples (key, header, width) defining displayed columns
        filename: Output filename
        translator: Translation manager for localized text
        table_name: Name of the table/data being exported
        include_all_fields: If True, export all fields from data, not just displayed columns
        selected_columns: List of tuples (key, header) for columns to export. If None, exports all.
        
    Returns:
        True if successful, False otherwise
    """
    try:
        import openpyxl
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise ImportError("openpyxl library not installed. Please run: pip install openpyxl")
    
    try:
        # Get global print settings
        settings = get_print_settings() if get_print_settings else None
        
        # Create workbook
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = table_name[:31]  # Excel sheet name max 31 chars
        
        # Set RTL if Arabic
        is_rtl = translator.current_language == 'ar'
        if is_rtl:
            ws.sheet_view.rightToLeft = True
        
        current_row = 1
        
        # Add professional header section
        current_row = add_excel_header(ws, settings, table_name, current_row, is_rtl)
        
        if not data:
            wb.save(filename)
            return True
        
        # Determine export columns
        if selected_columns:
            export_columns = selected_columns
        elif include_all_fields:
            all_keys = set()
            for row in data:
                all_keys.update(row.keys())
            
            export_columns = []
            column_dict = {col[0]: col[1] for col in columns}
            displayed_keys = [col[0] for col in columns]
            other_keys = sorted([k for k in all_keys if k not in displayed_keys])
            
            for key in displayed_keys + other_keys:
                header = column_dict.get(key, format_field_name(key))
                export_columns.append((key, header))
        else:
            export_columns = [(col[0], col[1]) for col in columns]
        
        # Arabic font support
        arabic_font_name = get_arabic_font_name() if is_rtl else 'Segoe UI'
        
        # Style definitions with Arabic font support
        header_fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=10, name=arabic_font_name)
        alt_fill = PatternFill(start_color="F8F9FA", end_color="F8F9FA", fill_type="solid")
        border = Border(
            left=Side(style='thin', color='DEE2E6'),
            right=Side(style='thin', color='DEE2E6'),
            top=Side(style='thin', color='DEE2E6'),
            bottom=Side(style='thin', color='DEE2E6')
        )
        
        # Data font for Arabic
        data_font = Font(size=10, name=arabic_font_name)
        data_font_muted = Font(size=10, name=arabic_font_name, color='7F8C8D', italic=True)
        
        # Alignment for Arabic (right-aligned for RTL)
        header_align = Alignment(horizontal='center', vertical='center')
        data_align = Alignment(horizontal='right' if is_rtl else 'left', vertical='top', wrap_text=True)
        center_align = Alignment(horizontal='center', vertical='center')
        
        # Add row number header
        cell = ws.cell(row=current_row, column=1, value='#')
        cell.font = Font(bold=True, color="FFFFFF", size=10, name='Segoe UI')  # Numbers stay LTR
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = border
        ws.column_dimensions['A'].width = 6
        
        # Add data column headers
        for col_idx, (_, col_name) in enumerate(export_columns, start=2):
            cell = ws.cell(row=current_row, column=col_idx, value=str(col_name))
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = border
        
        current_row += 1
        
        # Add data rows
        for row_idx, row_data in enumerate(data, start=1):
            # Row number
            cell = ws.cell(row=current_row, column=1, value=row_idx)
            cell.alignment = center_align
            cell.font = Font(bold=True, color='7F8C8D', name='Segoe UI')  # Numbers stay LTR
            cell.border = border
            if row_idx % 2 == 0:
                cell.fill = alt_fill
            
            # Data columns
            for col_idx, (col_key, _) in enumerate(export_columns, start=2):
                value = row_data.get(col_key)
                display_value = format_value_for_excel(value, col_key)
                cell = ws.cell(row=current_row, column=col_idx, value=display_value)
                cell.alignment = data_align
                cell.border = border
                
                # Apply appropriate font
                if display_value == '-':
                    cell.font = data_font_muted
                else:
                    cell.font = data_font
                
                if row_idx % 2 == 0:
                    cell.fill = alt_fill
            
            current_row += 1
        
        # Set column widths (skip row number column)
        # Use provided column_widths if available, otherwise auto-adjust
        column_widths = column_widths or {}
        
        # Calculate total percentage to normalize widths if needed
        total_percentage = sum(column_widths.values()) if column_widths else 0
        num_columns = len(export_columns)
        
        # If we have width specifications, use them; otherwise auto-adjust
        has_width_specs = len(column_widths) > 0
        
        if has_width_specs:
            logger.info(f"Applying column widths - Total: {total_percentage}%, Columns: {num_columns}, Specs: {len(column_widths)}")
        
        for col_idx, (col_key, col_header) in enumerate(export_columns, start=2):
            column_letter = get_column_letter(col_idx)
            
            # Use provided width ratio if available
            if col_key in column_widths:
                width_percentage = column_widths[col_key]
                original_percentage = width_percentage
                
                # Normalize if total doesn't equal 100% (but only if we have specs for all columns)
                if has_width_specs and total_percentage != 100 and len(column_widths) == num_columns:
                    # Normalize to 100%
                    width_percentage = (width_percentage / total_percentage) * 100 if total_percentage > 0 else 100 / num_columns
                    logger.debug(f"Normalized {col_key} from {original_percentage}% to {width_percentage}%")
                
                # Convert percentage to character width
                # Excel column width is measured in characters (0-255)
                # A typical Excel sheet with default font can fit ~80-100 characters across
                # We'll use a more direct conversion: percentage of a reasonable total width
                # For better visibility, we'll use a range of 8-50 characters per column
                min_char_width = 8
                max_char_width = 50
                
                # Direct percentage-based conversion
                # If 100% total width = ~100 chars, then 1% = 1 char (roughly)
                # But we want to scale it better for visibility
                char_width = min_char_width + (width_percentage / 100.0) * (max_char_width - min_char_width)
                char_width = max(min_char_width, min(max_char_width, char_width))
                
                logger.debug(f"Setting {col_key} ({col_header}) to {char_width:.1f} chars ({width_percentage}%)")
                ws.column_dimensions[column_letter].width = char_width
            else:
                # Auto-adjust based on content
                max_length = len(str(col_header))
                
                # Check data lengths (sample first 100 rows)
                for row_data in data[:100]:
                    value = row_data.get(col_key)
                    display_value = format_value_for_excel(value, col_key)
                    if display_value:
                        # Handle multi-line values
                        lines = str(display_value).split('\n')
                        for line in lines:
                            max_length = max(max_length, len(line))
                
                # Set width with padding, cap at 50
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column_letter].width = max(adjusted_width, 10)
        
        # Add footer
        current_row += 1
        add_excel_footer(ws, settings, len(data), current_row, is_rtl, translator)
        
        # Freeze header rows
        ws.freeze_panes = f'A{current_row - len(data) - 1}'
        
        # Save workbook
        wb.save(filename)
        return True
    except Exception as e:
        raise Exception(f"Error exporting to Excel: {str(e)}")


def add_excel_header(ws, settings, table_name: str, start_row: int, is_rtl: bool = False) -> int:
    """Add professional header to Excel worksheet with Arabic support"""
    from openpyxl.styles import Font, Alignment
    
    current_row = start_row
    arabic_font = get_arabic_font_name() if is_rtl else 'Segoe UI'
    
    # Alignment based on RTL
    text_align = 'right' if is_rtl else 'left'
    opposite_align = 'left' if is_rtl else 'right'
    
    # Header line 1 (Organization)
    if settings and settings.header_left_line1:
        cell = ws.cell(row=current_row, column=1, value=settings.header_left_line1)
        cell.font = Font(bold=True, size=14, color='2C3E50', name=arabic_font)
        cell.alignment = Alignment(horizontal=text_align, vertical='center')
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=5)
        current_row += 1
    
    # Header line 2 (Department)
    if settings and settings.header_left_line2:
        cell = ws.cell(row=current_row, column=1, value=settings.header_left_line2)
        cell.font = Font(size=11, color='34495E', name=arabic_font)
        cell.alignment = Alignment(horizontal=text_align, vertical='center')
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=5)
        current_row += 1
    
    # Document number and date on right side (stays LTR for numbers)
    if settings:
        doc_number = settings.generate_doc_number()
        date_str = settings.get_formatted_date() if settings.header_right_show_date else ""
        
        if doc_number or date_str:
            info_text = f"#{doc_number}" if doc_number else ""
            if date_str:
                info_text += f"  |  {date_str}" if info_text else date_str
            
            cell = ws.cell(row=start_row, column=8, value=info_text)
            cell.font = Font(bold=True, size=11, color='2C3E50', name='Segoe UI')  # Numbers stay LTR
            cell.alignment = Alignment(horizontal=opposite_align)
    
    # Report title
    current_row += 1
    cell = ws.cell(row=current_row, column=1, value=table_name)
    cell.font = Font(bold=True, size=12, color='667EEA', name=arabic_font)
    cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=5)
    
    # Total records
    current_row += 1
    
    # Empty row before table
    current_row += 1
    
    return current_row


def add_excel_footer(ws, settings, total_records: int, start_row: int, is_rtl: bool = False, translator=None):
    """Add footer to Excel worksheet with Arabic support"""
    from openpyxl.styles import Font, Alignment
    
    arabic_font = get_arabic_font_name() if is_rtl else 'Segoe UI'
    
    parts = []
    
    if settings:
        if settings.footer_left:
            parts.append(settings.footer_left)
        if settings.footer_center:
            parts.append(settings.footer_center)
    
    # Localized total text using translator
    if translator:
        total_label = translator.tr('report_total')
        records_label = translator.tr('report_records')
        parts.append(f"{total_label}: {total_records} {records_label}")
    elif is_rtl:
        parts.append(f"الإجمالي: {total_records} سجلات")
    else:
        parts.append(f"Total: {total_records} records")
    
    if settings and settings.show_print_date:
        parts.append(datetime.now().strftime('%Y-%m-%d %H:%M'))
    
    footer_text = ' | '.join(parts)
    
    cell = ws.cell(row=start_row, column=1, value=footer_text)
    cell.font = Font(size=9, color='7F8C8D', italic=True, name=arabic_font)
    cell.alignment = Alignment(horizontal='center')
    ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=8)


def format_field_name(key: str) -> str:
    """Format field name for display (convert snake_case to Title Case)"""
    return key.replace('_', ' ').title()


def format_value_for_excel(value, col_key: str = None) -> str:
    """Format value for Excel export - handles None values properly"""
    if value is None:
        return '-'  # Display dash instead of empty or "None"
    elif isinstance(value, datetime):
        return value.strftime('%Y-%m-%d %H:%M:%S')
    elif isinstance(value, bool):
        return 'Yes' if value else 'No'
    elif isinstance(value, (int, float)) and col_key:
        # Check if this is importance column
        if 'importance' in col_key.lower():
            try:
                return f"{float(value) * 100:.0f}%"
            except:
                return str(value)
        elif col_key == 'record_type':
            type_map = {
                'analysis': 'Analysis',
                'content': 'Content',
                'source': 'Source'
            }
            return type_map.get(str(value).lower(), str(value))
        return str(value)
    
    return str(value) if value else '-'


def export_timeline_to_excel(events: List[Dict], file_path: str, translator, 
                             density_mode: str = 'comfortable', include_header: bool = True,
                             preserve_format: bool = True):
    """Export timeline events to Excel with preserved formatting"""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
    from widgets.timeline_widget import TimelineDesignSystem
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Timeline Events"
    
    settings = get_print_settings() if get_print_settings else None
    is_rtl = translator.current_language == 'ar' if translator else False
    arabic_font = get_arabic_font_name() if is_rtl else 'Segoe UI'
    
    # Add header if requested and settings available
    if include_header and settings:
        start_row = add_excel_header(ws, settings, "Timeline Events", 1, is_rtl)
    else:
        start_row = 1
    
    # Title
    title_cell = ws.cell(row=start_row, column=1, value="Timeline Events Report")
    title_cell.font = Font(size=18, bold=True, name=arabic_font, color='2C3E50')
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=8)
    start_row += 2
    
    # Headers
    headers = ['Date', 'Title', 'Description', 'People', 'Places', 'Classification', 'Source', 'Type']
    header_fill = PatternFill(start_color='3498DB', end_color='3498DB', fill_type='solid')
    header_font = Font(size=11, bold=True, color='FFFFFF', name=arabic_font)
    
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    
    start_row += 1
    
    # Get density config
    config = TimelineDesignSystem.get_density_config(density_mode)
    
    # Add events
    for event in events:
        # Get event date
        event_date = None
        date_fields = ['date_content', 'date_analysis', 'date_entry', 
                      'date_creation', 'source_date_entry', 'content_date_creation']
        for field in date_fields:
            if field in event and event[field]:
                date_val = event[field]
                if isinstance(date_val, datetime):
                    event_date = date_val
                elif isinstance(date_val, str):
                    try:
                        for fmt in ['%Y-%m-%d %H:%M:%S', '%Y-%m-%d', '%Y-%m-%dT%H:%M:%S']:
                            try:
                                event_date = datetime.strptime(date_val[:19], fmt)
                                break
                            except:
                                continue
                    except:
                        pass
                if event_date:
                    break
        
        date_str = event_date.strftime('%Y-%m-%d %H:%M') if event_date else 'No Date'
        
        # Get title
        title = ""
        if 'content_title' in event and event['content_title']:
            title = str(event['content_title'])
        elif 'classification' in event and event['classification']:
            title = str(event['classification'])
        elif 'source_name' in event and event['source_name']:
            title = str(event['source_name'])
        else:
            title = translator.tr('timeline_event', default='Event') if translator else 'Event'
        
        # Get description
        description = ""
        if 'content_data' in event and event['content_data']:
            content = str(event['content_data'])
            description = content[:500] + "..." if len(content) > 500 else content
        
        # Get metadata
        people = format_value_for_excel(event.get('list_names_people', ''))
        places = format_value_for_excel(event.get('list_names_places', ''))
        classification = format_value_for_excel(event.get('classification', ''))
        source = format_value_for_excel(event.get('source_name', ''))
        record_type = format_value_for_excel(event.get('record_type', ''))
        
        # Write row
        row_data = [date_str, title, description, people, places, classification, source, record_type]
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=start_row, column=col_idx, value=value)
            cell.font = Font(size=10, name=arabic_font)
            cell.alignment = Alignment(horizontal='left' if not is_rtl else 'right', 
                                    vertical='top', wrap_text=True)
        
        start_row += 1
    
    # Adjust column widths
    ws.column_dimensions['A'].width = 18  # Date
    ws.column_dimensions['B'].width = 30  # Title
    ws.column_dimensions['C'].width = 50  # Description
    ws.column_dimensions['D'].width = 25  # People
    ws.column_dimensions['E'].width = 25  # Places
    ws.column_dimensions['F'].width = 20  # Classification
    ws.column_dimensions['G'].width = 25  # Source
    ws.column_dimensions['H'].width = 15  # Type
    
    # Add footer
    if settings:
        add_excel_footer(ws, settings, len(events), start_row + 1, is_rtl, translator)
    
    # Save workbook
    wb.save(file_path)